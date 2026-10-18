"""
Document parser module for OOXML Word (.docx) and PDF test papers.
Extracts sections, question blocks, options, answer keys, and formatting metadata.
"""
import io
import re
import copy
from typing import List, Dict, Tuple, Optional, Any
import docx
from lxml import etree

from backend.models import (
    QuestionType,
    OptionLayoutType,
    OptionItem,
    QuestionBlockData,
    SectionData,
    DocumentAnalysisResult,
)

NS = {
    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
}

# Regex patterns
RE_QUESTION = re.compile(
    r'^(?:Câu|CÂU|Question|Bài)\s*([0-9]{1,3})\s*(?:\(([0-9,.]+)\s*(?:điểm|đ)\))?[\s.:)-]\s*(.*)',
    re.IGNORECASE
)
RE_SECTION = re.compile(
    r'^(?:PHẦN|Phần|PART)\s*([0-9IVXLCDM]+|[A-Z])[\s.:\-]?\s*(.*)',
    re.IGNORECASE
)
RE_OPT_LETTER = re.compile(r'^([A-D])[\s.:)-]\s*(.*)')
RE_SUB_LETTER = re.compile(r'^([a-d])[\s.:)-]\s*(.*)', re.IGNORECASE)
RE_END_MARKER = re.compile(r'---+\s*(?:HẾT|Hết|THE END)\s*---+', re.IGNORECASE)
RE_CODE_PLACEHOLDER = re.compile(r'(?:Mã\s*đề|MÃ\s*ĐỀ|Đề\s*số)[\s:]*([0-9A-Za-z]+)?', re.IGNORECASE)


class DocxParser:
    """Parser specifically designed for Vietnamese K-12 Word test papers."""

    def __init__(self, docx_bytes: bytes, filename: str = "document.docx"):
        self.docx_bytes = docx_bytes
        self.filename = filename
        self.doc = docx.Document(io.BytesIO(docx_bytes))
        self.body = self.doc._body._element

    def analyze(self) -> DocumentAnalysisResult:
        """Performs full analysis of the document without altering any XML nodes."""
        # 1. Detect answers if document has an answer table or grading guide
        detected_answers, grading_guide_found = self._extract_grading_answers()
        
        # 2. Check for student answer sheet table (e.g. Table with Câu 1..Câu N)
        has_student_answer_table = self._detect_student_answer_table()
        
        # 3. Check for code placeholder (e.g. MÃ ĐỀ: 000)
        code_found, code_text = self._detect_code_placeholder()

        # 4. Find end-of-exam index (where ---HẾT--- or section break begins)
        exam_end_idx = self._find_exam_end_index()

        # 5. Extract sections & questions
        sections, questions = self._extract_sections_and_questions(exam_end_idx, detected_answers)

        # Count types
        num_mc = sum(1 for q in questions if q.question_type == QuestionType.MULTIPLE_CHOICE)
        num_tf = sum(1 for q in questions if q.question_type == QuestionType.TRUE_FALSE)
        num_essay = sum(1 for q in questions if q.question_type == QuestionType.ESSAY)

        return DocumentAnalysisResult(
            filename=self.filename,
            file_type="docx",
            is_scanned_pdf=False,
            pdf_warning=None,
            total_questions=len(questions),
            num_mc_questions=num_mc,
            num_tf_questions=num_tf,
            num_essay_questions=num_essay,
            sections=sections,
            questions=questions,
            detected_answers=detected_answers,
            has_student_answer_table=has_student_answer_table,
            has_grading_guide=grading_guide_found,
            code_placeholder_found=code_found,
            code_placeholder_text=code_text,
            exam_end_body_idx=exam_end_idx,
        )

    def _find_exam_end_index(self) -> int:
        """Finds the index in body where the exam content ends and grading/guide begins."""
        total = len(self.body)
        for idx, child in enumerate(self.body):
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                txt = child.xpath('string(.)').strip()
                if RE_END_MARKER.search(txt):
                    return idx
                # Check for section break
                sectPr = child.xpath('./w:pPr/w:sectPr')
                if sectPr:
                    # If this paragraph has a section break and follows questions, likely end of exam
                    if idx > 10:
                        return idx
        return total

    def _extract_grading_answers(self) -> Tuple[Dict[int, str], bool]:
        """Extracts answer keys from grading tables at the end of the document if present."""
        answers: Dict[int, str] = {}
        found = False

        # Check all tables, especially tables near the end
        for t_idx, table in enumerate(self.doc.tables):
            # Check for nested tables inside cells (very common in Hướng dẫn chấm)
            for row in table.rows:
                for cell in row.cells:
                    for nested in cell.tables:
                        ans_dict = self._parse_nested_answer_table(nested)
                        if ans_dict:
                            answers.update(ans_dict)
                            found = True

            # Also check direct table if it's an answer table
            ans_dict = self._parse_direct_answer_table(table)
            if ans_dict:
                answers.update(ans_dict)
                found = True

        return answers, found

    def _parse_nested_answer_table(self, table: docx.table.Table) -> Dict[int, str]:
        """Parses a table formatted like:
        Row 0: 1, 2, 3... 14
        Row 1: C, B, B... A
        Row 2: 15, 16... 28
        Row 3: C, D... Sai
        """
        ans_dict = {}
        if len(table.rows) >= 2 and len(table.columns) >= 3:
            # Check row 0 for question numbers
            try:
                # Row 0 & 1
                row0 = [c.text.strip() for c in table.rows[0].cells]
                row1 = [c.text.strip() for c in table.rows[1].cells]
                if row0[0].isdigit() and (int(row0[0]) == 1):
                    for q_str, a_str in zip(row0, row1):
                        if q_str.isdigit():
                            ans_dict[int(q_str)] = a_str
                
                # Row 2 & 3 if present
                if len(table.rows) >= 4:
                    row2 = [c.text.strip() for c in table.rows[2].cells]
                    row3 = [c.text.strip() for c in table.rows[3].cells]
                    if row2[0].isdigit():
                        for q_str, a_str in zip(row2, row3):
                            if q_str.isdigit():
                                ans_dict[int(q_str)] = a_str
            except Exception:
                pass
        return ans_dict

    def _parse_direct_answer_table(self, table: docx.table.Table) -> Dict[int, str]:
        """Parses 2-column or grid answer table: [Câu, Đáp án] or similar."""
        ans_dict = {}
        try:
            # Pattern: Col 0 is 'Câu 1', Col 1 is 'A'
            if len(table.columns) == 2:
                for row in table.rows:
                    c0 = row.cells[0].text.strip()
                    c1 = row.cells[1].text.strip()
                    m = re.match(r'^(?:Câu\s*)?([0-9]{1,3})', c0, re.IGNORECASE)
                    if m and c1:
                        ans_dict[int(m.group(1))] = c1
        except Exception:
            pass
        return ans_dict

    def _detect_student_answer_table(self) -> bool:
        """Detects if document has a student answer sheet table (e.g. Table 2 with Câu 1..Câu 28)."""
        for table in self.doc.tables:
            for row in table.rows:
                row_text = ' '.join(c.text.strip() for c in row.cells)
                if 'Câu 1' in row_text and 'Câu 2' in row_text:
                    return True
        return False

    def _detect_code_placeholder(self) -> Tuple[bool, Optional[str]]:
        """Checks for existing Mã đề string in paragraphs or tables."""
        for p in self.doc.paragraphs:
            m = RE_CODE_PLACEHOLDER.search(p.text)
            if m:
                return True, m.group(0)
        for t in self.doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    m = RE_CODE_PLACEHOLDER.search(cell.text)
                    if m:
                        return True, m.group(0)
        return False, None

    def _extract_sections_and_questions(
        self, exam_end_idx: int, detected_answers: Dict[int, str]
    ) -> Tuple[List[SectionData], List[QuestionBlockData]]:
        sections: List[SectionData] = []
        questions: List[QuestionBlockData] = []

        # First pass: locate all section headings and question lead paragraphs
        items_map = []  # list of (body_idx, 'section'|'question', data)
        for idx in range(exam_end_idx):
            child = self.body[idx]
            tag = child.tag.split('}')[-1]
            if tag == 'p':
                txt = child.xpath('string(.)').strip()
                # Check section
                sec_m = RE_SECTION.match(txt)
                if sec_m:
                    items_map.append((idx, 'section', txt))
                    continue
                # Check question
                q_m = RE_QUESTION.match(txt)
                if q_m:
                    q_num = int(q_m.group(1))
                    lead_txt = txt
                    items_map.append((idx, 'question', (q_num, lead_txt)))

        # If no explicit sections found, create a default section
        current_section_id = 1
        current_section = SectionData(
            id=1,
            title="ĐỀ THI TỔNG HỢP",
            heading_body_idx=0,
            question_ids=[],
            shuffle_questions=True,
            shuffle_options=True,
        )
        sections.append(current_section)

        # Build slices
        for i, (idx, item_type, data) in enumerate(items_map):
            if item_type == 'section':
                # Start new section
                current_section_id = len(sections) + 1
                current_section = SectionData(
                    id=current_section_id,
                    title=data,
                    heading_body_idx=idx,
                    question_ids=[],
                    # Default: disable question shuffle for Tự luận if specified in title
                    shuffle_questions='tự luận' not in data.lower(),
                    shuffle_options='đúng' not in data.lower() and 'tự luận' not in data.lower(),
                )
                sections.append(current_section)

            elif item_type == 'question':
                q_num, lead_txt = data
                # End of this question block is the start of next item, or exam_end_idx
                next_idx = items_map[i + 1][0] if i + 1 < len(items_map) else exam_end_idx
                block_elems = [self.body[k] for k in range(idx, next_idx)]

                # Analyze question structure
                q_data = self._analyze_question_block(
                    q_num=q_num,
                    lead_txt=lead_txt,
                    block_elems=block_elems,
                    start_idx=idx,
                    end_idx=next_idx,
                    section_id=current_section_id,
                    section_title=current_section.title,
                    correct_answer=detected_answers.get(q_num, ""),
                )
                questions.append(q_data)
                current_section.question_ids.append(q_data.id)

        # Remove default section if empty and other sections exist
        if len(sections) > 1 and len(sections[0].question_ids) == 0:
            sections = sections[1:]

        # Normalize section IDs to 1, 2, 3...
        sec_id_map = {}
        for new_id, sec in enumerate(sections, start=1):
            sec_id_map[sec.id] = new_id
            sec.id = new_id

        for q in questions:
            if q.section_id in sec_id_map:
                q.section_id = sec_id_map[q.section_id]

        return sections, questions

    def _analyze_question_block(
        self,
        q_num: int,
        lead_txt: str,
        block_elems: List[Any],
        start_idx: int,
        end_idx: int,
        section_id: int,
        section_title: str,
        correct_answer: str,
    ) -> QuestionBlockData:
        q_id = f"q_{q_num}"
        has_tables = any(e.tag.endswith('tbl') for e in block_elems)
        has_images = any(bool(e.xpath('.//w:drawing | .//w:pict')) for e in block_elems)
        has_equations = any(bool(e.xpath('.//m:oMath')) for e in block_elems)

        # Determine option layout and extract options
        option_layout = OptionLayoutType.NONE
        q_type = QuestionType.ESSAY
        options: List[OptionItem] = []
        is_essay_section = 'tự luận' in section_title.lower() or 'tự luận' in lead_txt.lower()

        # 1. Check if block has a 2x2 table for options
        if len(block_elems) >= 2 and block_elems[1].tag.endswith('tbl'):
            tbl = docx.table.Table(block_elems[1], self.doc)
            if len(tbl.rows) == 2 and len(tbl.columns) == 2:
                # Check if cells start with A, B, C, D
                cells = [tbl.rows[0].cells[0], tbl.rows[0].cells[1], tbl.rows[1].cells[0], tbl.rows[1].cells[1]]
                c_texts = [c.text.strip() for c in cells]
                if all(RE_OPT_LETTER.match(ct) for ct in c_texts):
                    option_layout = OptionLayoutType.TABLE_2X2
                    q_type = QuestionType.MULTIPLE_CHOICE
                    letters = ['A', 'B', 'C', 'D']
                    for j, ct in enumerate(c_texts):
                        m = RE_OPT_LETTER.match(ct)
                        opt_body = m.group(2) if m else ct
                        options.append(OptionItem(
                            letter=letters[j],
                            text=opt_body,
                            full_text=ct,
                            original_index=j,
                        ))

        # 2. Check if options are 4 paragraphs
        if option_layout == OptionLayoutType.NONE and len(block_elems) >= 5:
            p_texts = [block_elems[k].xpath('string(.)').strip() for k in range(1, 5)]
            if all(RE_OPT_LETTER.match(pt) for pt in p_texts):
                option_layout = OptionLayoutType.PARAGRAPHS_4
                q_type = QuestionType.MULTIPLE_CHOICE
                letters = ['A', 'B', 'C', 'D']
                for j, pt in enumerate(p_texts):
                    m = RE_OPT_LETTER.match(pt)
                    opt_body = m.group(2) if m else pt
                    options.append(OptionItem(
                        letter=letters[j],
                        text=opt_body,
                        full_text=pt,
                        original_index=j,
                    ))

        # 3. Check for True/False (ý a, b, c, d hoặc nhận định Đúng/Sai)
        if option_layout == OptionLayoutType.NONE:
            if is_essay_section:
                q_type = QuestionType.ESSAY
            else:
                # Check if block has sub-items a), b), c), d)
                p_sub_texts = [e.xpath('string(.)').strip() for e in block_elems[1:] if e.tag.endswith('p')]
                has_sub_letters = any(RE_SUB_LETTER.match(pt) for pt in p_sub_texts)
                if has_sub_letters or correct_answer in ['Đúng', 'Sai', 'True', 'False']:
                    q_type = QuestionType.TRUE_FALSE
                elif 'đúng hoặc sai' in lead_txt.lower() or 'chọn đúng hoặc sai' in lead_txt.lower():
                    q_type = QuestionType.TRUE_FALSE
                elif any(lead_txt.rstrip().endswith(ch) for ch in ['…', '...', '.....']):
                    # Sentence with dots evaluation
                    if correct_answer in ['Đúng', 'Sai']:
                        q_type = QuestionType.TRUE_FALSE
                    else:
                        q_type = QuestionType.ESSAY
                else:
                    q_type = QuestionType.ESSAY

        # Build full text preview
        preview_parts = []
        for e in block_elems:
            if e.tag.endswith('p'):
                preview_parts.append(e.xpath('string(.)').strip())
            elif e.tag.endswith('tbl'):
                t_obj = docx.table.Table(e, self.doc)
                t_txt = ' | '.join(c.text.strip().replace('\n', ' ') for r in t_obj.rows for c in r.cells)
                preview_parts.append(f"[Bảng: {t_txt[:60]}...]")
        full_preview = '\n'.join(p for p in preview_parts if p)

        return QuestionBlockData(
            id=q_id,
            original_number=q_num,
            lead_text=lead_txt,
            full_preview=full_preview,
            question_type=q_type,
            option_layout=option_layout,
            options=options,
            correct_answer=correct_answer,
            is_locked=False,
            # For True/False, default shuffle_options to False per Requirement 5
            shuffle_options=(q_type == QuestionType.MULTIPLE_CHOICE),
            section_id=section_id,
            has_images=has_images,
            has_tables=has_tables,
            has_equations=has_equations,
            start_body_idx=start_idx,
            end_body_idx=end_idx,
        )
